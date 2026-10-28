"""Export chat answers (Markdown) to Word, Excel and Markdown files."""
import io
import re
from datetime import datetime

import pandas as pd
from docx import Document
from docx.shared import Pt, RGBColor

RED = RGBColor(0xE6, 0x00, 0x00)
BLACK = RGBColor(0x1C, 0x1C, 0x1C)
TABLE_ROW = re.compile(r"^\s*\|.*\|\s*$")
TABLE_SEP = re.compile(r"^\s*\|?\s*:?-{2,}:?\s*(\|\s*:?-{2,}:?\s*)*\|?\s*$")


# ---------- Markdown tables ----------
def _cells(line: str) -> list[str]:
    return [c.strip() for c in line.strip().strip("|").split("|")]


def extract_tables(md: str) -> list[pd.DataFrame]:
    tables, lines, i = [], md.splitlines(), 0
    while i < len(lines) - 1:
        if TABLE_ROW.match(lines[i]) and TABLE_SEP.match(lines[i + 1]):
            header, rows, i = _cells(lines[i]), [], i + 2
            while i < len(lines) and TABLE_ROW.match(lines[i]):
                rows.append(_cells(lines[i]))
                i += 1
            rows = [(r + [""] * len(header))[:len(header)] for r in rows]
            tables.append(pd.DataFrame(rows, columns=header))
        else:
            i += 1
    return tables


def tables_to_xlsx(md: str) -> bytes:
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as w:
        for n, df in enumerate(extract_tables(md), 1):
            df.to_excel(w, sheet_name=f"Table {n}", index=False)
    return buf.getvalue()


# ---------- Markdown -> Word ----------
def _runs(par, text: str):
    for part in re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text):
        if part.startswith("**") and part.endswith("**"):
            par.add_run(part[2:-2]).bold = True
        elif part.startswith("`") and part.endswith("`"):
            par.add_run(part[1:-1]).font.name = "Consolas"
        elif part:
            par.add_run(part)


def _write(doc: Document, md: str):
    lines, i = md.splitlines(), 0
    while i < len(lines):
        line, s = lines[i], lines[i].strip()
        if TABLE_ROW.match(line) and i + 1 < len(lines) and TABLE_SEP.match(lines[i + 1]):
            header, rows, i = _cells(line), [], i + 2
            while i < len(lines) and TABLE_ROW.match(lines[i]):
                rows.append(_cells(lines[i]))
                i += 1
            t = doc.add_table(rows=1, cols=len(header))
            t.style = "Light Grid Accent 1"
            for c, h in enumerate(header):
                t.rows[0].cells[c].text = h
            for r in rows:
                cells = t.add_row().cells
                for c in range(len(header)):
                    cells[c].text = r[c] if c < len(r) else ""
            continue
        m = re.match(r"^(#{1,6})\s+(.*)", s)
        if m:
            doc.add_heading(m.group(2), level=min(len(m.group(1)), 3))
        elif re.match(r"^[-*•]\s+", s):
            _runs(doc.add_paragraph(style="List Bullet"), re.sub(r"^[-*•]\s+", "", s))
        elif re.match(r"^\d+[.)]\s+", s):
            _runs(doc.add_paragraph(style="List Number"), re.sub(r"^\d+[.)]\s+", "", s))
        elif s:
            _runs(doc.add_paragraph(), s)
        i += 1


def _doc(title: str) -> Document:
    doc = Document()
    doc.styles["Normal"].font.size = Pt(10.5)
    for r in doc.add_heading(title, level=0).runs:
        r.font.color.rgb = BLACK
    p = doc.add_paragraph()
    tag = p.add_run("RISK ATLAS  ")
    tag.bold, tag.font.color.rgb, tag.font.size = True, RED, Pt(8)
    p.add_run(f"Exported {datetime.now():%b %d, %Y %H:%M}").font.size = Pt(8)
    return doc


def _bytes(doc: Document) -> bytes:
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def answer_to_docx(md: str) -> bytes:
    doc = _doc("Risk Atlas Answer")
    _write(doc, md)
    return _bytes(doc)


def chat_to_docx(messages: list[dict]) -> bytes:
    doc = _doc("Risk Atlas Conversation")
    for m in messages:
        run = doc.add_paragraph().add_run("You" if m["role"] == "user" else "Risk Atlas")
        run.bold, run.font.color.rgb = True, RED if m["role"] == "assistant" else BLACK
        _write(doc, m["content"])
    return _bytes(doc)


def chat_to_md(messages: list[dict]) -> str:
    out = [f"# Risk Atlas Conversation ({datetime.now():%b %d, %Y %H:%M})\n"]
    for m in messages:
        out.append(f"**{'You' if m['role'] == 'user' else 'Risk Atlas'}**\n\n{m['content']}\n")
    return "\n".join(out)
